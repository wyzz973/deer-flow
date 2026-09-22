"""Word export: a report someone can hand to a reader.

The Markdown and HTML exports already give a reader clickable citations, the
date a source states and a note when only a search excerpt backs a claim. Word
was the one export that dropped all three and printed Mermaid diagrams as
source code. These tests pin the repaired contract, including the pictures the
browser rendered.
"""

from __future__ import annotations

import base64
from io import BytesIO

import pytest
from docx import Document

from deepresearch import report as documents

# A real 1x1 PNG: the smallest input that still has to pass the byte check.
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")

FLOWCHART = "flowchart TD\n    A[开始] --> B[结束]"

# The page a report prints on, less its margins.
TEXT_WIDTH = 6.5
TEXT_HEIGHT = 8.4


def page(eid, url="https://docs.example.org/guide", **extra):
    return {
        "evidence_id": eid,
        "url": url,
        "canonical_url": url,
        "source_uri": None,
        "source_name": "web",
        "origin": "external",
        "title": "官方指南",
        "snippet": "原文摘录",
        "provenance": "fetched_document",
        "document_hash": "snapshot",
        "source_level": "L2",
        **extra,
    }


def build(body, pool=None, *, summary="摘要 [[E001]]", demo=False, **extra):
    pool = pool or {"E001": page("E001")}
    document = documents.assemble("研究报告", summary, [("技术分析", body)], lang="zh")
    mapping = documents.bind(document, pool)
    value = {
        "format": "markdown-v2",
        "document": document,
        "citation_map": mapping,
        "citations": documents.citations(mapping, pool),
        "demo": demo,
        "language": "zh",
        **extra,
    }
    return value


def export(value, diagrams=None):
    return Document(BytesIO(documents.docx_document(value, diagrams=diagrams)))


def text_of(doc):
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def test_a_citation_in_the_body_is_a_superscript_link_to_its_reference():
    doc = export(build("结论成立。[[E001]]"))
    xml = doc.element.xml
    assert 'w:anchor="ref-1"' in xml, "正文的 [1] 必须能跳到文末条目"
    assert 'w:name="ref-1"' in xml, "参考资料条目必须有对应书签"
    assert "superscript" in xml, "引用编号应是上标，不是正文里的方括号"


def test_a_reference_entry_carries_the_link_the_site_and_the_stated_date():
    pool = {"E001": page("E001", published_at="2026-07-10T00:00:00Z")}
    doc = export(build("结论成立。[[E001]]", pool))
    entry = next(p.text for p in doc.paragraphs if p.text.startswith("1."))
    assert "官方指南" in entry
    assert "docs.example.org" in entry, "读者要能一眼看出这条来自哪个站点"
    assert "2026-07-10" in entry, "发布日期决定时效，Markdown 导出有，Word 不能丢"
    targets = [rel.target_ref for rel in doc.part.rels.values() if rel.is_external]
    assert "https://docs.example.org/guide" in targets, "标题要是真正可点击的超链接"


def test_a_reference_backed_only_by_a_search_excerpt_says_so():
    pool = {"E001": page("E001", provenance="observed_source", snippet="检索片段")}
    doc = export(build("结论成立。[[E001]]", pool))
    entry = next(p.text for p in doc.paragraphs if p.text.startswith("1."))
    assert "检索摘录" in entry, "读者必须能分辨哪条读过原文、哪条只有检索片段"


def test_a_rendered_mermaid_diagram_becomes_a_picture_with_a_caption():
    body = f"下图说明流程。\n\n```mermaid\n{FLOWCHART}\n```"
    key = documents.diagram_key(FLOWCHART)
    doc = export(build(body), diagrams={key: PNG})
    assert "图 1" in text_of(doc), "图要有题注，读者才能在正文里指代它"
    assert "flowchart TD" not in text_of(doc), "渲染成功后正文不该再出现源码"
    assert len(doc.inline_shapes) == 1, "图应当作为图片嵌入"


def test_an_unrendered_diagram_degrades_to_a_caption_and_an_appendix():
    body = f"下图说明流程。\n\n```mermaid\n{FLOWCHART}\n```"
    doc = export(build(body))
    whole = text_of(doc)
    assert "图 1" in whole
    assert doc.inline_shapes == [] or len(doc.inline_shapes) == 0
    appendix = whole.split("附录")[-1]
    assert "flowchart TD" in appendix, "源码不能丢，但要收进附录"
    assert "flowchart TD" not in whole.split("附录")[0], "正文里不能再倒源码"


def test_only_real_png_bytes_are_embedded():
    body = f"```mermaid\n{FLOWCHART}\n```"
    key = documents.diagram_key(FLOWCHART)
    doc = export(build(body), diagrams={key: b"<svg>not a png</svg>"})
    assert len(doc.inline_shapes) == 0, "伪装成图片的字节不能进入要分发的文档"
    assert "图 1" in text_of(doc), "被拒绝的图仍要留下题注和降级说明"


def test_an_ordinary_code_block_is_still_a_code_block():
    doc = export(build("示例：\n\n```python\nprint(1)\n```"))
    assert "print(1)" in text_of(doc), "非 Mermaid 的代码块保持原样"


def test_a_table_reads_like_a_report_table_not_a_spreadsheet():
    """Three rules and no vertical lines, the shape a published report uses.

    A full grid with a shaded header is loud on paper and, where Word breaks
    the table, the boxes make the seam obvious.
    """
    body = "| 维度 | 方案 |\n| --- | --- |\n| 并发 | **单写者** [[E001]] |"
    doc = export(build(body))
    table = doc.tables[0]
    header = table.rows[0]
    borders = table._tbl.xml
    assert all(run.bold for cell in header.cells for p in cell.paragraphs for run in p.runs), "表头要加粗"
    assert "tblHeader" in header._tr.xml, "跨页时表头要重复"
    assert 'w:insideV w:val="none"' in borders, "不要竖线"
    assert 'w:insideH w:val="none"' in borders, "数据行之间不要横线"
    assert "w:shd" not in borders, "不要底纹"


def test_emphasis_inside_a_table_cell_survives():
    body = "| 维度 | 方案 |\n| --- | --- |\n| 并发 | **单写者** [[E001]] |"
    doc = export(build(body))
    cell = doc.tables[0].rows[1].cells[1]
    assert any(run.bold for p in cell.paragraphs for run in p.runs), "单元格里的加粗以前被正则删掉了"
    assert cell.text.startswith("单写者"), "强调标记本身不该出现在文字里"


def test_the_document_opens_with_a_clickable_table_of_contents():
    doc = export(build("正文。[[E001]]"))
    whole = text_of(doc)
    assert "目录" in whole
    assert whole.index("目录") < whole.index("技术分析"), "目录在正文之前"
    assert 'w:anchor="heading-' in doc.element.xml, "目录条目要能跳转"


def test_every_page_is_numbered():
    doc = export(build("正文。[[E001]]"))
    footer = doc.sections[0].footer
    assert "PAGE" in footer.paragraphs[0]._p.xml, "对外发的报告要有页码"


def test_chinese_headings_get_an_east_asian_font():
    doc = export(build("正文。[[E001]]"))
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        assert style.element.xml.count("eastAsia") >= 1, f"{name} 缺中文字体回退，Word 里会掉字形"


def test_the_demo_notice_comes_before_the_report_not_after_it():
    doc = export(build("正文。[[E001]]", demo=True))
    whole = text_of(doc)
    assert "演示模式" in whole
    assert whole.index("演示模式") < whole.index("技术分析"), "警示放在末尾等于没放"


def test_a_report_stored_before_dates_and_basis_existed_still_exports():
    value = build("正文。[[E001]]")
    value["citations"] = [{k: v for k, v in ref.items() if k not in ("basis", "published_at", "domain")} for ref in value["citations"]]
    doc = export(value)
    assert "官方指南" in text_of(doc)


@pytest.mark.parametrize("size", [1, 3])
def test_diagrams_are_numbered_in_document_order(size):
    sources = [f"flowchart TD\n    A{index} --> B{index}" for index in range(size)]
    body = "\n\n".join(f"```mermaid\n{source}\n```" for source in sources)
    diagrams = {documents.diagram_key(source): PNG for source in sources}
    doc = export(build(body), diagrams=diagrams)
    whole = text_of(doc)
    assert [f"图 {index + 1}" in whole for index in range(size)] == [True] * size
    assert len(doc.inline_shapes) == size


@pytest.mark.asyncio
async def test_the_export_endpoint_embeds_the_browser_rendered_diagrams(settings, monkeypatch):
    """The picture comes from the reader's browser; the report never does."""
    import sys
    import types

    import httpx
    from fastapi import FastAPI

    from deepresearch.api import build_router
    from deepresearch.service import ResearchService

    body = f"流程如下。\n\n```mermaid\n{FLOWCHART}\n```"
    value = {**build(body), "version": 1}
    service = ResearchService(settings)
    await service.store.start()
    await service.store.create({"run_id": "run-1", "owner": "alice", "status": "COMPLETED", "report": value}, "k", "h")
    def resolver(request):
        user = request.headers.get("x-test-user")
        return types.SimpleNamespace(user_id=user) if user else None

    monkeypatch.setitem(sys.modules, "deerflow_extension_api", types.SimpleNamespace(resolve_principal=resolver))
    app = FastAPI()
    app.include_router(build_router(service))
    key = documents.diagram_key(FLOWCHART)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://localhost") as client:
        headers = {"x-test-user": "alice"}
        sent = {"diagrams": {key: base64.b64encode(PNG).decode()}}
        assert (await client.post("/api/deepresearch/run-1/report/docx", json=sent)).status_code == 401
        assert (await client.post("/api/deepresearch/run-1/report/docx", json=sent, headers={"x-test-user": "bob"})).status_code == 404
        response = await client.post("/api/deepresearch/run-1/report/docx", json=sent, headers=headers)
        assert response.status_code == 200 and response.headers["content-disposition"].endswith('.docx"')
        assert len(Document(BytesIO(response.content)).inline_shapes) == 1
        # The same run with no pictures still exports, and still has no code in it.
        plain = await client.get("/api/deepresearch/run-1/report?format=docx", headers=headers)
        rendered = Document(BytesIO(plain.content))
        assert len(rendered.inline_shapes) == 0 and "图 1" in text_of(rendered)


def png_of(width, height):
    """A real PNG of a given size; its shape is what the export has to respect."""
    import struct
    import zlib

    def chunk(kind, payload):
        body = kind + payload
        return struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0)
    raw = b"".join(b"\x00" + b"\xff" * width for _ in range(height))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


@pytest.mark.parametrize(
    ("pixels", "shape"),
    [
        ((2416, 534), "wide"),  # a left-to-right flowchart
        ((562, 1868), "tall"),  # a top-down one: 19 inches at its own scale
        ((948, 1840), "tall"),
    ],
)
def test_a_diagram_is_never_larger_than_the_page_it_prints_on(pixels, shape):
    """A picture Word cannot fit on a page is a picture the reader loses."""
    width, height = pixels
    body = f"```mermaid\n{FLOWCHART}\n```"
    doc = export(build(body), diagrams={documents.diagram_key(FLOWCHART): png_of(width, height)})
    picture = doc.inline_shapes[0]
    assert picture.width.inches <= TEXT_WIDTH + 0.01, "不能超出正文栏宽"
    assert picture.height.inches <= TEXT_HEIGHT + 0.01, f"{shape} 图高 {picture.height.inches:.2f} 英寸，一页放不下会被裁掉"
    # The diagram must still be the shape it was drawn in.
    assert abs(picture.width.inches / picture.height.inches - width / height) < 0.01


def test_a_table_row_never_breaks_across_a_page():
    """Word splits a row by default, which strands a line of a cell alone on
    the next page under a repeated header. A row moves whole instead."""
    body = "| 维度 | 说明 |\n| --- | --- |\n| 并发 | 很长的一段说明文字，足以让这一行跨过分页处 |"
    doc = export(build(body))
    rows = doc.tables[0].rows
    assert all("cantSplit" in row._tr.xml for row in rows), "每一行都不能被分页切开"


def test_table_columns_have_declared_widths():
    """Word's autofit gives one column most of the width and squeezes the rest;
    a report table reads better with declared, even columns."""
    body = "| 维度 | 方案 A | 方案 B |\n| --- | --- | --- |\n| 并发 | 单写者 | 多写者 |"
    doc = export(build(body))
    table = doc.tables[0]
    assert table.autofit is False
    assert "fixed" in table._tbl.xml, "表格要用固定布局"
    widths = [cell.width.inches for cell in table.rows[0].cells]
    assert all(width > 0 for width in widths)
    assert abs(sum(widths) - TEXT_WIDTH) < 0.05, f"列宽合计应当填满正文栏，实为 {sum(widths):.2f} 英寸"


def test_the_page_gives_the_report_its_full_width():
    doc = export(build("正文。[[E001]]"))
    section = doc.sections[0]
    usable = (section.page_width - section.left_margin - section.right_margin) / 914400
    assert abs(usable - TEXT_WIDTH) < 0.02, f"正文栏宽 {usable:.2f} 英寸"


def test_a_very_tall_diagram_gets_a_page_of_its_own():
    """Fitting a 1:3 diagram to the page height leaves it a narrow ribbon, and
    dropping it mid-text strands half of the previous page. Given its own page
    it is as large as the page allows and nothing is left ragged."""
    body = f"前面一段文字。[[E001]]\n\n```mermaid\n{FLOWCHART}\n```"
    doc = export(build(body), diagrams={documents.diagram_key(FLOWCHART): png_of(562, 1868)})
    holder = next(p for p in doc.paragraphs if "graphic" in p._p.xml)
    assert "pageBreakBefore" in holder._p.xml, "超高的图要另起一页"
    assert "keepNext" in holder._p.xml, "题注不能和图分开"


def test_a_diagram_that_fits_stays_in_the_flow():
    body = f"前面一段文字。[[E001]]\n\n```mermaid\n{FLOWCHART}\n```"
    doc = export(build(body), diagrams={documents.diagram_key(FLOWCHART): png_of(2416, 534)})
    holder = next(p for p in doc.paragraphs if "graphic" in p._p.xml)
    assert "pageBreakBefore" not in holder._p.xml, "放得下的图不该赶到下一页"
