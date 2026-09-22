from io import BytesIO

import pytest
from pydantic import ValidationError

from deepresearch import report as documents
from deepresearch.contracts import ResearchPlan, ResearchUnit, StructuredReport, safe_http_url
from deepresearch.evidence import canonical_url, merge_results, ordered_sources
from deepresearch.render import citation_metadata, docx_report
from deepresearch.validators import research_gaps, single_source, supplemental_units


def make_document(ids, extra=""):
    marker = "[[" + ", ".join(ids) + "]]" if ids else ""
    summary = f"**可验证结论** {marker}"
    return documents.assemble("报告", summary, [("技术分析", f"可验证结论。{marker}{extra}")], [], [], "zh")


def test_plan_dag_and_unique_ids(plan):
    with pytest.raises(ValidationError):
        ResearchPlan(goal="valid goal", research_units=[ResearchUnit(id="R1", skill="technical-route", objective="a good objective", depends_on=["R1"])])
    with pytest.raises(ValidationError):
        ResearchPlan(goal="valid goal", research_units=plan.research_units * 2)
    with pytest.raises(ValidationError):
        ResearchPlan(goal="valid goal", research_units=[ResearchUnit(id="R1", skill="technical-route", objective="a good objective", depends_on=["missing"])])


@pytest.mark.parametrize("url", ["javascript:alert(1)", "file:///etc/passwd", "https://user:pass@example.org", "https://example.org/\nsecret", "https://example.org:notaport/path"])
def test_url_deny(url):
    with pytest.raises(ValueError):
        safe_http_url(url)


def test_canonical_url_preserves_business_parameters():
    assert canonical_url("https://EXAMPLE.com:443/page?id=1&utm_source=abc&lang=zh#x") == "https://example.com/page?id=1&lang=zh"
    assert canonical_url("http://example.com:8080/") == "http://example.com:8080/"
    assert canonical_url("http://example.com:0/") == "http://example.com:0/"


def test_merge_replay_lineage_and_distinct_origins(result):
    pool, findings, lineage = merge_results([result])
    replay, repeated, _ = merge_results([result], pool)
    assert pool == replay and findings == repeated
    assert set(pool) == {"E001", "E002"}
    assert set(lineage.values()) == set(pool)
    second = result.model_copy(deep=True)
    second.unit_id = "R2"
    merged, _, _ = merge_results([result, second], pool)
    assert len(merged) == 2
    assert all(e["unit_ids"] == ["R1", "R2"] for e in merged.values())


def test_display_metadata_from_the_web_is_bounded_not_rejected(result):
    """A page without a title is labeled by its URL, and URLs can be enormous.

    Truncating display text keeps one monster link from failing a whole research
    run in evidence_merge; identity fields stay strict.
    """
    from deepresearch.contracts import RawEvidence

    url = "https://example.com/image.png?jwt=" + "e" * 1600
    evidence = RawEvidence(raw_id="raw-long", title=url, url=url, origin="external", source_name="external-web", publisher="p" * 400, snippet="正文")
    assert len(evidence.title) == 1000 and evidence.title.endswith("…") and evidence.title.startswith("https://example.com/image.png")
    assert len(evidence.publisher) == 200
    assert evidence.url == url  # the locator itself is never silently altered
    # An over-long excerpt is bounded; empty text has nothing to cite and is refused.
    assert len(RawEvidence(raw_id="raw-long", title="t", url="https://example.com/a", origin="external", source_name="external-web", publisher="p", snippet="正" * 30000).snippet) == 20000
    with pytest.raises(ValidationError):
        RawEvidence(raw_id="raw-empty", title="t", url="https://example.com/a", origin="external", source_name="external-web", publisher="p", snippet="   ")
    # A result carrying such evidence still merges.
    body = result.model_dump(mode="json")
    body["raw_evidences"].append({**evidence.model_dump(mode="json"), "raw_id": "raw-long"})
    from deepresearch.contracts import ResearchResult

    revalidated = ResearchResult.model_validate(body)
    pool, _findings, _lineage = merge_results([revalidated], {})
    assert any(str(item["title"]).startswith("https://example.com/image.png") for item in pool.values())


def test_source_priority_injected_file_fallback(settings, plan, tmp_path):
    priority = tmp_path / "priority.yaml"
    priority.write_text("- external-web\n", encoding="utf-8")
    settings.source_priority_file = str(priority)
    a, provenance = ordered_sources(settings, plan.research_units[0].source_strategy, ["internal-kb"])
    assert a[0].name == "internal-kb" and provenance == "injected"
    b, provenance = ordered_sources(settings, plan.research_units[0].source_strategy)
    assert b[0].name == "external-web" and provenance == "file"
    settings.source_priority_file = None
    _, provenance = ordered_sources(settings, plan.research_units[0].source_strategy)
    assert provenance == "fallback"
    assert {s.origin for s in b} == {"internal", "external"}


def test_reference_binding_first_appearance_and_render(result):
    pool, _, _ = merge_results([result])
    document = make_document(["E002", "E001"])
    mapping = documents.bind(document, pool)
    assert mapping == {"E002": 1, "E001": 2}
    assert not documents.validate(document, pool, set(pool))
    exported = documents.export_markdown(document, mapping, pool)
    assert exported == documents.export_markdown(document, mapping, pool)
    assert exported.count("[1](#ref-1)[2](#ref-2)") == 2
    assert "id=1" in exported and "utm_source=x" in exported  # Export the original locator, not a rewritten signed URL.
    assert all("utm_source" not in (e["canonical_url"] or "") for e in pool.values())
    display = documents.display_markdown(document, mapping)
    assert "[1](#citation-E002)[2](#citation-E001)" in display and "[[" not in display
    with pytest.raises(ValueError):
        documents.bind(make_document(["E999"]), pool)


def test_adjacent_markers_merge_and_code_is_never_rewritten(result):
    pool, _, _ = merge_results([result])
    body = "结论。[[E001]] [[E001, E002]]\n\n```mermaid\nflowchart LR\n  A[[E999]] --> B\n```\n\n`list[[E999]]` 是类型示例。"
    document = documents.assemble("报告", "摘要 [[E001]]", [("分析", body)], lang="zh")
    assert documents.marker_ids(document) == ["E001", "E001", "E001", "E002"]
    mapping = documents.bind(document, pool)
    display = documents.display_markdown(document, mapping)
    assert "结论。[1](#citation-E001)[2](#citation-E002)" in display
    assert "A[[E999]] --> B" in display and "`list[[E999]]`" in display
    assert not documents.problems(document, set(pool))


def test_html_escapes_model_and_source_text(result):
    pool, _, _ = merge_results([result])
    document = make_document(["E001"], extra="\n\n<script>alert(1)</script>")
    pool["E001"]["title"] = '<img src=x onerror="alert(1)">'
    value = documents.html_document(document, {"E001": 1}, pool)
    assert "<script>" not in value and "<img " not in value
    assert '<sup><a href="#ref-1">[1]</a></sup>' in value


def test_writer_problems_are_precise_and_sanitizing_never_invents_citations(result):
    pool, _, _ = merge_results([result])
    draft = "- 有据结论。[[E001]]\n- 无据结论。[[E999]] 另一句仍有据。[[E002]]\n| 维度 | A |\n| --- | --- |\n| 并发 | 伪造 [[E998]] |\n见 [官网](https://example.org/doc) 与 https://example.org/raw，编号[3]。"
    errors = documents.problems(draft, set(pool))
    assert any("E998, E999" in error for error in errors)
    assert any("numeric citation" in error for error in errors)
    assert any("2 URL" in error for error in errors)
    cleaned, audit = documents.sanitize(draft, set(pool))
    assert "E999" not in cleaned and "E998" not in cleaned
    assert "- 有据结论。[[E001]]" in cleaned and "另一句仍有据。[[E002]]" in cleaned
    assert "| 并发 | — |" in cleaned
    assert "https://" not in cleaned and "官网" in cleaned and "[3]" not in cleaned
    assert audit == {"dropped_statements": 2, "stripped_links": 2, "stripped_numeric_citations": 1}
    assert not documents.problems(cleaned, set(pool))


def test_final_gate_requires_citations_and_a_title(result):
    pool, _, _ = merge_results([result])
    assert any("cites no evidence" in error for error in documents.validate(make_document([]), pool, set(pool)))
    assert not documents.validate(make_document([]), pool, set())
    assert any("no title" in error for error in documents.validate("## 只有章节 [[E001]]", pool, set(pool)))


def test_section_answers_are_normalized_without_touching_statements():
    answer = "```markdown\n## 并发写入\n\n## 小节\n正文。[[E001]]\n```"
    assert documents.clean_answer(answer, "并发写入") == "### 小节\n正文。[[E001]]"
    assert documents.clean_answer("# 报告\n\n## 章节", whole_document=True) == "# 报告\n\n## 章节"
    document = documents.assemble("标题", "摘要", [("第一章", "### 小节\n正文")], ["百人规模"], ["公开资料有限"], "zh")
    assert documents.toc(document) == [{"level": 2, "text": "执行摘要"}, {"level": 2, "text": "第一章"}, {"level": 3, "text": "小节"}, {"level": 2, "text": "研究范围与局限"}]
    assert "- 百人规模" in document and "- 公开资料有限" in document
    assert documents.title_of(document) == "标题"


def test_writer_preambles_are_dropped_but_real_openings_are_kept():
    body = "**判断**。事实。[[E001]]"
    assert documents.clean_answer("Writing the assigned section body now.\n\n" + body, "章节") == body
    assert documents.clean_answer("好的，下面开始撰写本节正文。\n\n" + body, "章节") == body
    table = "以下是三种路径的对比：\n\n| 路径 | 成本 |\n|---|---|\n| A | 低 [[E001]] |"
    assert documents.clean_answer(table, "章节") == table
    assert documents.clean_answer("Here is the revised report.\n\n# 报告\n\n## 章节", whole_document=True) == "# 报告\n\n## 章节"


def test_notes_about_how_the_text_was_produced_are_removed():
    body = "**结论**。事实。[[E001]]\n\n*说明：本摘要为对各章节草稿的综合，未引入草稿之外的事实，也未新增证据 ID。*"
    assert documents.clean_answer(body, "执行摘要") == "**结论**。事实。[[E001]]"
    kept = "- 研究对象：证据 ID 在取证系统中的作用 [[E001]]\n\n| 项 | 说明 |\n|---|---|\n| A | evidence ID [[E001]] |"
    assert documents.clean_answer(kept, "章节") == kept
    code = '```mermaid\nflowchart LR\n  A["证据 ID"] --> B\n```'
    assert documents.clean_answer(code, "章节") == code


def test_outline_headings_lose_numbering_and_generated_parts_are_recognized():
    assert documents.plain_heading("一、前沿模型代际盘点") == "前沿模型代际盘点"
    assert documents.plain_heading("3. Pricing and packaging") == "Pricing and packaging"
    assert documents.plain_heading("第二章 市场格局") == "市场格局"
    assert documents.plain_heading("（四）风险与治理") == "风险与治理"
    for heading in ["2026 年路线图", "3.5 Sonnet 与 GPT-5 对比", "CI/CD 成熟度"]:
        assert documents.plain_heading(heading) == heading
    for heading in ["执行摘要与读者行动清单", "一、执行摘要", "Executive Summary", "摘要", "参考来源", "研究范围与局限", "四、研究限制与未解决的证据缺口", "References"]:
        assert documents.reserved_heading(heading), heading
    for heading in ["结论与行动建议", "风险与局限", "Sources of competitive advantage", "摘要生成能力对比"]:
        assert not documents.reserved_heading(heading), heading


def test_citable_gaps_ignore_private_context_and_hedge_single_sites(result, plan):
    pool, findings, _ = merge_results([result])
    units = [u.model_dump() for u in plan.research_units]
    assert research_gaps(plan, units, findings, pool) == []
    only = [dict(findings[0], evidence_ids=["E001"], high_risk=True)]
    gaps = research_gaps(plan, units, only, pool)
    assert any(g["code"].startswith("missing-") for g in gaps)
    assert not any(g["code"] == "independence" for g in gaps)
    assert single_source(only[0], pool, set(pool))
    assert not single_source(dict(findings[0], high_risk=True), pool, set(pool))
    # Non-citable evidence does not count as coverage.
    assert any(g["code"] == "coverage" for g in research_gaps(plan, units, findings, pool, citable_ids=set()))
    context = result.model_dump()
    context["assumptions_needed"] = ["团队规模未知"]
    assert research_gaps(plan, units, findings, pool, [context]) == []
    supplemental = supplemental_units(plan, gaps, 1, 1)
    assert len(supplemental) == 1 and supplemental[0]["parent_gap_id"]
    assert supplemental[0]["depends_on"] == ["R1"]


def test_supplements_spend_budget_on_missing_coverage_first():
    plan = ResearchPlan(goal="valid goal", research_units=[ResearchUnit(id=uid, skill="technical-route", objective="a good objective") for uid in ("R1", "R2")])
    gaps = [{"gap_id": "R1-open-questions", "unit_id": "R1", "code": "open-questions", "description": "q"}, {"gap_id": "R2-coverage", "unit_id": "R2", "code": "coverage", "description": "c"}]
    selected = supplemental_units(plan, gaps, 1, 1)
    assert [unit["depends_on"] for unit in selected] == [["R2"]]
    assert supplemental_units(plan, gaps, 1, 0) == []


def test_docx_export_uses_the_verified_markdown_document(result):
    from docx import Document

    pool, _, _ = merge_results([result])
    body = "正文结论。[[E001]]\n\n| 维度 | 方案 A |\n| --- | --- |\n| 并发 | 单写者 [[E002]] |\n\n- **要点** 说明"
    document = documents.assemble("报告", "摘要 [[E001]]", [("技术分析", body)], lang="zh")
    mapping = documents.bind(document, pool)
    value = {"format": "markdown-v2", "document": document, "citation_map": mapping, "citations": documents.citations(mapping, pool), "demo": True}
    doc = Document(BytesIO(documents.docx_document(value)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "正文结论。[1]" in text and "演示模式" in text and "参考资料" in text
    assert doc.tables[0].rows[1].cells[1].text == "单写者 [2]"
    assert "要点 说明" in text


def test_historical_ast_reports_still_export_to_word(result):
    from docx import Document

    pool, _, _ = merge_results([result])
    segment = {"text": "可验证结论", "evidence_ids": ["E001"], "segment_type": "fact"}
    report = StructuredReport.model_validate({"title": "报告", "executive_summary": [segment], "sections": [{"heading": "技术分析", "unit_ids": ["R1"], "segments": [segment]}], "conclusion": [segment]})
    value = {"report": report.model_dump(), "citation_map": {"E001": 1}, "citations": citation_metadata({"E001": 1}, pool), "demo": True, "limitations": []}
    doc = Document(BytesIO(docx_report(value)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "可验证结论[1]" in text and "演示模式" in text and "参考资料" in text


def test_a_single_unusable_evidence_record_never_fails_the_whole_unit():
    """Evidence is derived from open-web payloads and may predate a schema rule.

    evidence_merge revalidates every stored result, so one record the current
    contract rejects used to fail the entire research. Such records are dropped
    with their reason, findings keep the references that survive, and a finding
    left without any reference goes with them.
    """
    from deepresearch.evidence import valid_result

    def record(raw_id, **overrides):
        return {
            "raw_id": raw_id,
            "title": "可用来源",
            "url": "https://example.com/" + raw_id,
            "origin": "external",
            "source_name": "external-web",
            "source_level": "L2",
            "publisher": "example",
            "snippet": "正文片段",
            "provenance": "fetched_document",
            **overrides,
        }

    body = {
        "unit_id": "u1",
        "confidence": 0.6,
        "raw_evidences": [
            record("keep1"),
            record("longtitle", title="https://example.com/x?jwt=" + "e" * 1800),
            record("credentials", url="https://user:secret@example.com/page"),
            record("emptytext", snippet="   "),
            record("keep2"),
        ],
        "findings": [
            {"claim": "结论一", "raw_evidence_refs": ["keep1", "emptytext"], "confidence": 0.7},
            {"claim": "只靠被丢弃的证据", "raw_evidence_refs": ["credentials"], "confidence": 0.5},
        ],
    }
    result, dropped = valid_result(body)
    # A long title is display metadata and is truncated, not dropped.
    kept = [item.raw_id for item in result.raw_evidences]
    assert kept == ["keep1", "longtitle", "keep2"]
    assert [item["raw_id"] for item in dropped] == ["credentials", "emptytext"]
    assert {field for item in dropped for field in item["fields"]} == {"url", "snippet"}
    assert [finding.raw_evidence_refs for finding in result.findings] == [["keep1"]]
    assert result.findings[0].claim == "结论一"


def test_a_structurally_broken_result_still_fails_loudly():
    from pydantic import ValidationError

    from deepresearch.evidence import valid_result

    with pytest.raises(ValidationError):
        valid_result({"unit_id": "不是标识符 ", "confidence": 2.5, "raw_evidences": [], "findings": []})


def test_display_labels_a_model_overshoots_are_bounded_not_rejected():
    """A step label that runs a few characters long must not throw away a plan.

    Titles are labels for the plan card; claims and briefs stay strict so the
    conversion retry can give the model precise feedback instead of silently
    cutting research content.
    """
    plan = ResearchPlan.model_validate(
        {
            "goal": "比较方案",
            "title": "标" * 400,
            "research_units": [{"id": "u1", "skill": "technical-route", "title": "查" * 300, "objective": "对比运维成本"}],
        }
    )
    assert len(plan.title) == 120 and plan.title.endswith("…")
    assert len(plan.research_units[0].title) == 80
    with pytest.raises(ValidationError):  # research content is never silently cut
        ResearchPlan.model_validate({"goal": "比" * 20000, "research_units": [{"id": "u1", "skill": "s", "objective": "o"}]})


def test_a_date_requirement_is_checked_against_the_dates_sources_actually_stated(plan, result):
    """A caller-supplied ``not_before`` is worth checking once evidence has dates.

    The check used to require every evidence to be provenance ``document``,
    which research never produces, so a date requirement was silently only a
    model judgement. Now that a source's stated date is kept, a cutoff can be
    enforced where it is provable: dated evidence, all of it too old. Undated
    evidence stays a matter for the researcher — supplements cannot make a
    provider send dates it does not have.
    """
    from datetime import UTC, datetime

    pool, findings, _ = merge_results([result])
    units = [u.model_dump() for u in plan.research_units]
    plan.research_units[0].source_strategy.not_before = datetime(2026, 1, 1, tzinfo=UTC)
    codes = lambda: {g["code"] for g in research_gaps(plan, units, findings, pool)}  # noqa: E731 - three one-line calls below
    # Nothing is dated: unchanged behaviour, the researcher reports the doubt.
    assert "date" not in codes()
    for item in pool.values():
        item["published_at"] = "2019-05-04T00:00:00+00:00"
    assert "date" in codes()
    # One source inside the window satisfies it, even if the others are old.
    next(iter(pool.values()))["published_at"] = "2026-03-09T00:00:00+00:00"
    assert "date" not in codes()
