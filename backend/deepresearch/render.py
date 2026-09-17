"""Citation metadata and exports for historical structured (AST) reports.

New reports are Markdown documents handled by ``report.py``. Reports published
before that change keep their stored Markdown/HTML, and this module still
exports them to Word from the same verified AST. Source metadata is read-only.
"""

from __future__ import annotations

from .contracts import StructuredReport


def citation_metadata(mapping, pool):
    groups = {}
    for eid, number in mapping.items():
        if number not in groups:
            groups[number] = {"number": number, **pool[eid], "evidence_ids": [], "excerpts": []}
        groups[number]["evidence_ids"].append(eid)
        groups[number]["excerpts"].append({"evidence_id": eid, "text": pool[eid]["snippet"][:2000], "document_hash": pool[eid].get("document_hash")})
    return sorted(groups.values(), key=lambda item: item["number"])


def docx_report(value):
    """Export a historical AST report; no Markdown reparsing or model formatting."""
    from io import BytesIO

    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    report = StructuredReport.model_validate(value["report"])
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.add_heading(report.title, 0)
    if value["demo"]:
        doc.add_paragraph("演示模式：合成测试数据，不可用于业务决策。")
    mapping = value["citation_map"]

    def segment_text(item):
        return item.text + "".join(f"[{number}]" for number in dict.fromkeys(mapping[e] for e in item.evidence_ids))

    def segments(items):
        for item in items:
            doc.add_paragraph(segment_text(item))

    doc.add_heading("执行摘要", 1)
    segments(report.executive_summary)
    if report.comparison_table:
        source = report.comparison_table
        doc.add_heading("快速对比", 1)
        table = doc.add_table(rows=1, cols=len(source.headers) + 1)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ["维度", *source.headers], strict=True):
            cell.text = label
        for row in source.rows:
            for cell, text in zip(table.add_row().cells, [row.label, *[segment_text(item) for item in row.cells]], strict=True):
                cell.text = text
    for section in report.sections:
        doc.add_heading(section.heading, 1)
        segments(section.segments)
    doc.add_heading("结论", 1)
    segments(report.conclusion)
    if value["limitations"]:
        doc.add_heading("研究限制", 1)
        for text in value["limitations"]:
            doc.add_paragraph(text)
    doc.add_heading("参考资料", 1)
    for ref in value["citations"]:
        doc.add_paragraph(f"[{ref['number']}] {ref['title']}\n{ref.get('url') or ref['canonical_url'] or ref['source_uri']}")
    # Stable core metadata avoids exposing local usernames.
    doc.core_properties.author = "DeepResearch"
    doc.core_properties.title = report.title
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
